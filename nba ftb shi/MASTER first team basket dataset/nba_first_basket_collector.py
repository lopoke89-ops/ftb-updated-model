"""
NBA First Basket Dataset Collector - 2025-26 Season
=====================================================
Pulls play-by-play data from NBA.com for every completed game in the
2025-26 regular season and extracts:
  - First basket scorer for each team
  - Shot type, time of first basket, quarter, score type (2pt/3pt)

Run once to build the full dataset, then re-run nightly after games finish
to append new results. The script is smart about skipping already-collected
game IDs so repeated runs are fast.

Requirements:
    pip install nba_api openpyxl pandas

Usage:
    python nba_first_basket_collector.py
    python nba_first_basket_collector.py --output my_file.xlsx
"""

import argparse
import time
import pandas as pd
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from pathlib import Path

try:
    from nba_api.stats.endpoints import leaguegamelog, playbyplayv3
    from nba_api.stats.static import teams as nba_teams_static
except ImportError:
    raise SystemExit("Missing dependency: pip install nba_api")

OUTPUT_FILE = "NBA_First_Basket_2025-26.xlsx"
SEASON = "2025-26"
DELAY = 1.5       # seconds between successful API calls
MAX_RETRIES = 5   # retry attempts on timeout/rate-limit
RETRY_BACKOFF = [5, 15, 30, 60, 120]  # wait (s) before each retry

# ─── Helpers ──────────────────────────────────────────────────────────────────

TEAM_ID_MAP  = {t["id"]: t["abbreviation"] for t in nba_teams_static.get_teams()}
TEAM_FULL_MAP = {t["id"]: t["full_name"]   for t in nba_teams_static.get_teams()}
TEAM_ABB_MAP  = {t["abbreviation"]: t["id"] for t in nba_teams_static.get_teams()}


def get_all_completed_games() -> pd.DataFrame:
    """Return a DataFrame of all completed 2025-26 regular season games."""
    print("Fetching season game log...")
    gl = leaguegamelog.LeagueGameLog(
        season=SEASON,
        season_type_all_star="Regular Season",
        player_or_team_abbreviation="T",   # team-level log
    )
    df = gl.get_data_frames()[0]
    time.sleep(DELAY)

    # Each game appears twice (one row per team). De-duplicate to unique GAME_IDs.
    games = (
        df[["GAME_ID", "GAME_DATE", "MATCHUP", "TEAM_ID", "TEAM_ABBREVIATION"]]
        .copy()
    )

    # Build home/away context from MATCHUP string ("LAL vs. GSW" = home, "LAL @ GSW" = away)
    games["IS_HOME"] = games["MATCHUP"].str.contains(" vs\\. ")
    home = games[games["IS_HOME"]].rename(columns={
        "TEAM_ID": "HOME_TEAM_ID",
        "TEAM_ABBREVIATION": "HOME_TEAM",
    })
    away = games[~games["IS_HOME"]].rename(columns={
        "TEAM_ID": "AWAY_TEAM_ID",
        "TEAM_ABBREVIATION": "AWAY_TEAM",
    })

    merged = home.merge(
        away[["GAME_ID", "AWAY_TEAM_ID", "AWAY_TEAM"]],
        on="GAME_ID",
    )[["GAME_ID", "GAME_DATE", "HOME_TEAM_ID", "HOME_TEAM", "AWAY_TEAM_ID", "AWAY_TEAM"]]
    merged = merged.drop_duplicates("GAME_ID").sort_values("GAME_DATE").reset_index(drop=True)

    print(f"  Found {len(merged)} completed games.")
    return merged


def get_first_baskets(game_id: str, home_team_id: int, away_team_id: int) -> dict:
    """
    Pull play-by-play for a game (V3) and return the first made field goal per team.
    Handles both camelCase and UPPER_CASE column schemas that NBA.com returns.
    """
    import re

    plays = None
    for attempt in range(MAX_RETRIES):
        try:
            pbp = playbyplayv3.PlayByPlayV3(game_id=game_id)
            plays = pbp.get_data_frames()[0]
            time.sleep(DELAY)
            break  # success — exit retry loop
        except Exception as e:
            err = str(e)
            is_timeout = "timed out" in err.lower() or "timeout" in err.lower()
            is_ratelimit = "429" in err or "rate" in err.lower()
            if (is_timeout or is_ratelimit) and attempt < MAX_RETRIES - 1:
                wait = RETRY_BACKOFF[attempt]
                print(f"    TIMEOUT/RATE-LIMIT on game {game_id} (attempt {attempt+1}/{MAX_RETRIES}) — waiting {wait}s...")
                time.sleep(wait)
            else:
                print(f"    ERROR fetching PBP for game {game_id}: {e}")
                return {}
    if plays is None:
        return {}

    if plays.empty:
        print(f"    WARNING: empty play-by-play for game {game_id}")
        return {}

    cols = set(plays.columns.tolist())

    # Detect schema: camelCase (newer) vs UPPER_CASE (older nba_api builds)
    if "actionType" in cols:
        action_col   = "actionType"
        desc_col     = "description"
        team_col     = "teamId"
        player_col   = "playerNameI" if "playerNameI" in cols else "playerName"
        period_col   = "period"
        clock_col    = "clock"
        shot_res_col = "shotResult" if "shotResult" in cols else None
        sub_col      = "subType"    if "subType"    in cols else None
    elif "ACTION_TYPE" in cols:
        action_col   = "ACTION_TYPE"
        desc_col     = "DESCRIPTION"
        team_col     = "TEAM_ID"
        player_col   = "PLAYER_NAME"
        period_col   = "PERIOD"
        clock_col    = "CLOCK"
        shot_res_col = "SHOT_RESULT" if "SHOT_RESULT" in cols else None
        sub_col      = "SUB_TYPE"   if "SUB_TYPE"   in cols else None
    else:
        print(f"    WARNING: unknown PBP schema for {game_id}. Cols: {list(cols)[:15]}")
        return {}

    # Filter to made field goals
    made = plays[plays[action_col].str.contains("Made Shot|2pt Shot|3pt Shot", case=False, na=False)].copy()

    if made.empty and shot_res_col:
        shot_mask   = plays[action_col].str.contains("Shot", case=False, na=False)
        result_mask = plays[shot_res_col].str.lower() == "made"
        made = plays[shot_mask & result_mask].copy()

    if made.empty:
        print(f"    WARNING: no made shots found for game {game_id}")
        return {}

    result = {}
    for label, team_id in [("HOME", home_team_id), ("AWAY", away_team_id)]:
        team_plays = made[made[team_col] == team_id]
        if team_plays.empty:
            continue
        first = team_plays.iloc[0]
        description = str(first.get(desc_col, "")).strip()

        # Parse ISO 8601 clock "PT11M39.00S" -> "11:39"
        raw_clock = str(first.get(clock_col, ""))
        m = re.search(r"PT(\d+)M([\d.]+)S", raw_clock)
        clock_str = f"{int(m.group(1))}:{int(float(m.group(2))):02d}" if m else raw_clock

        # Determine 2pt vs 3pt
        sub = str(first.get(sub_col, "") if sub_col else "").upper()
        action_val = str(first.get(action_col, "")).upper()
        is_three = "3PT" in description.upper() or "3PT" in sub or "3PT" in action_val or "3-POINT" in sub
        points = 3 if is_three else 2

        result[label] = {
            "PLAYER_NAME" : str(first.get(player_col, "")),
            "TEAM_ID"     : team_id,
            "PERIOD"      : first.get(period_col, 1),
            "PCTIMESTRING": clock_str,
            "DESCRIPTION" : description,
            "POINTS"      : points,
        }

    return result


def load_existing(path: str) -> set:
    """Return set of GAME_IDs already in the Excel file (complete or not)."""
    p = Path(path)
    if not p.exists():
        return set()
    wb = load_workbook(p, read_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    if "GAME_ID" not in headers:
        return set()
    idx = headers.index("GAME_ID")
    ids = {row[idx].value for row in ws.iter_rows(min_row=2) if row[idx].value}
    wb.close()
    return ids


def load_incomplete_rows(path: str) -> list[dict]:
    """
    Scan the Excel file for rows where either team's first scorer is missing.
    Returns a list of dicts with keys: row_num, game_id, game_date,
    home_team, away_team, missing_home, missing_away.
    """
    p = Path(path)
    if not p.exists():
        return []
    wb = load_workbook(p, read_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]

    col = {h: i for i, h in enumerate(headers)}
    required = {"GAME_ID", "HOME_TEAM", "AWAY_TEAM", "HOME_FIRST_SCORER", "AWAY_FIRST_SCORER"}
    if not required.issubset(col):
        wb.close()
        return []

    incomplete = []
    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        def val(key):
            v = row[col[key]]
            return str(v).strip() if v is not None else ""

        home_scorer = val("HOME_FIRST_SCORER")
        away_scorer = val("AWAY_FIRST_SCORER")
        missing_home = home_scorer in ("", "N/A", "None")
        missing_away = away_scorer in ("", "N/A", "None")

        if missing_home or missing_away:
            incomplete.append({
                "row_num"     : row_num,
                "game_id"     : val("GAME_ID"),
                "game_date"   : val("GAME_DATE"),
                "home_team"   : val("HOME_TEAM"),
                "away_team"   : val("AWAY_TEAM"),
                "missing_home": missing_home,
                "missing_away": missing_away,
            })

    wb.close()
    return incomplete


# Column index map for the First Baskets sheet (1-based)
_FB_COLS = {
    "GAME_ID": 1, "GAME_DATE": 2, "HOME_TEAM": 3, "AWAY_TEAM": 4,
    "HOME_FIRST_SCORER": 5, "HOME_SHOT_DESC": 6, "HOME_POINTS": 7,
    "HOME_TIME": 8, "HOME_PERIOD": 9,
    "AWAY_FIRST_SCORER": 10, "AWAY_SHOT_DESC": 11, "AWAY_POINTS": 12,
    "AWAY_TIME": 13, "AWAY_PERIOD": 14,
}


def patch_excel_row(path: str, row_num: int, baskets: dict,
                    missing_home: bool, missing_away: bool):
    """Overwrite only the blank cells in an existing row with fresh data."""
    wb = load_workbook(path)
    ws = wb["First Baskets"]

    base_font = Font(name="Arial", size=10)
    center    = Alignment(horizontal="center")
    fill      = (PatternFill("solid", start_color="DCE6F1")
                 if row_num % 2 == 0
                 else PatternFill("solid", start_color="FFFFFF"))

    def write(col_name, value):
        c = ws.cell(row=row_num, column=_FB_COLS[col_name], value=value)
        c.font = base_font
        c.alignment = center
        c.fill = fill

    if missing_home and "HOME" in baskets:
        h = baskets["HOME"]
        write("HOME_FIRST_SCORER", h["PLAYER_NAME"])
        write("HOME_SHOT_DESC",    h["DESCRIPTION"])
        write("HOME_POINTS",       h["POINTS"])
        write("HOME_TIME",         h["PCTIMESTRING"])
        write("HOME_PERIOD",       h["PERIOD"])

    if missing_away and "AWAY" in baskets:
        a = baskets["AWAY"]
        write("AWAY_FIRST_SCORER", a["PLAYER_NAME"])
        write("AWAY_SHOT_DESC",    a["DESCRIPTION"])
        write("AWAY_POINTS",       a["POINTS"])
        write("AWAY_TIME",         a["PCTIMESTRING"])
        write("AWAY_PERIOD",       a["PERIOD"])

    wb.save(path)



def write_excel(records: list[dict], path: str):
    """Write / append records to a nicely formatted Excel file."""
    p = Path(path)

    if p.exists():
        wb = load_workbook(p)
        ws = wb.active
        start_row = ws.max_row + 1
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "First Baskets"
        start_row = 2

        # ── Header row ────────────────────────────────────────────────────
        headers = [
            "GAME_ID", "GAME_DATE", "HOME_TEAM", "AWAY_TEAM",
            "HOME_FIRST_SCORER", "HOME_SHOT_DESC", "HOME_POINTS", "HOME_TIME", "HOME_PERIOD",
            "AWAY_FIRST_SCORER", "AWAY_SHOT_DESC", "AWAY_POINTS", "AWAY_TIME", "AWAY_PERIOD",
        ]
        header_fill = PatternFill("solid", start_color="17375E")
        header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Column widths
        widths = [14, 12, 12, 12, 22, 40, 8, 8, 8, 22, 40, 8, 8, 8]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.row_dimensions[1].height = 30

    # Alternating row fill
    fill_a = PatternFill("solid", start_color="DCE6F1")
    fill_b = PatternFill("solid", start_color="FFFFFF")
    base_font = Font(name="Arial", size=10)
    center = Alignment(horizontal="center")

    for i, rec in enumerate(records):
        row = start_row + i
        fill = fill_a if row % 2 == 0 else fill_b
        values = [
            rec["GAME_ID"], rec["GAME_DATE"], rec["HOME_TEAM"], rec["AWAY_TEAM"],
            rec.get("HOME_PLAYER", "N/A"), rec.get("HOME_DESC", "N/A"),
            rec.get("HOME_PTS", ""), rec.get("HOME_TIME", ""), rec.get("HOME_PERIOD", ""),
            rec.get("AWAY_PLAYER", "N/A"), rec.get("AWAY_DESC", "N/A"),
            rec.get("AWAY_PTS", ""), rec.get("AWAY_TIME", ""), rec.get("AWAY_PERIOD", ""),
        ]
        for col_idx, val in enumerate(values, 1):
            cell = ws.cell(row=row, column=col_idx, value=val)
            cell.fill = fill
            cell.font = base_font
            cell.alignment = center

    # Freeze panes under header
    ws.freeze_panes = "A2"

    # Auto-filter
    ws.auto_filter.ref = ws.dimensions

    wb.save(path)
    print(f"  Saved {len(records)} new records → {path}")


# ─── Summary Sheet ────────────────────────────────────────────────────────────

def add_summary_sheet(path: str):
    """Add/refresh a Summary tab with player and team first-basket counts."""
    wb = load_workbook(path)
    if "Summary" in wb.sheetnames:
        del wb["Summary"]

    data_ws = wb["First Baskets"]
    rows = list(data_ws.iter_rows(min_row=2, values_only=True))

    player_counts: dict[str, int] = {}
    team_counts: dict[str, int] = {}

    for row in rows:
        # HOME_FIRST_SCORER at index 4, HOME_TEAM at 2
        # AWAY_FIRST_SCORER at index 9, AWAY_TEAM at 3
        for player_col, team_col in [(4, 2), (9, 3)]:
            player = row[player_col]
            team = row[team_col]
            if player and player != "N/A":
                player_counts[player] = player_counts.get(player, 0) + 1
            if team and team != "N/A":
                team_counts[team] = team_counts.get(team, 0) + 1

    ws = wb.create_sheet("Summary")

    title_font = Font(name="Arial", bold=True, size=12, color="FFFFFF")
    hdr_fill = PatternFill("solid", start_color="17375E")
    bold = Font(name="Arial", bold=True, size=10)
    normal = Font(name="Arial", size=10)
    center = Alignment(horizontal="center")

    # Player leaderboard
    ws["A1"] = "Player First Basket Leaderboard"
    ws["A1"].font = title_font
    ws["A1"].fill = hdr_fill
    ws.merge_cells("A1:B1")

    ws["A2"] = "Player"
    ws["B2"] = "# First Baskets"
    for c in ["A2", "B2"]:
        ws[c].font = bold
        ws[c].fill = PatternFill("solid", start_color="2E74B5")
        ws[c].font = Font(name="Arial", bold=True, color="FFFFFF")
        ws[c].alignment = center

    sorted_players = sorted(player_counts.items(), key=lambda x: -x[1])
    for r, (player, cnt) in enumerate(sorted_players, 3):
        ws.cell(row=r, column=1, value=player).font = normal
        cell = ws.cell(row=r, column=2, value=cnt)
        cell.font = normal
        cell.alignment = center

    # Team counts (column D-E)
    ws["D1"] = "Team First Basket Leaderboard"
    ws["D1"].font = title_font
    ws["D1"].fill = hdr_fill
    ws.merge_cells("D1:E1")

    ws["D2"] = "Team"
    ws["E2"] = "# First Baskets"
    for c in ["D2", "E2"]:
        ws[c].font = Font(name="Arial", bold=True, color="FFFFFF")
        ws[c].fill = PatternFill("solid", start_color="2E74B5")
        ws[c].alignment = center

    sorted_teams = sorted(team_counts.items(), key=lambda x: -x[1])
    for r, (team, cnt) in enumerate(sorted_teams, 3):
        ws.cell(row=r, column=4, value=team).font = normal
        cell = ws.cell(row=r, column=5, value=cnt)
        cell.font = normal
        cell.alignment = center

    for col in ["A", "B", "D", "E"]:
        ws.column_dimensions[col].width = 26

    ws.freeze_panes = "A3"

    wb.save(path)
    print("  Summary sheet updated.")


# ─── Main ─────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description="NBA First Basket Dataset Builder")
    parser.add_argument("--output", default=OUTPUT_FILE, help="Output Excel file path")
    args = parser.parse_args()

    output = args.output
    existing_ids = load_existing(output)
    print(f"Already collected: {len(existing_ids)} games.")

    games_df = get_all_completed_games()
    new_games = games_df[~games_df["GAME_ID"].isin(existing_ids)]
    print(f"New games to process: {len(new_games)}")

    if not new_games.empty:
        new_records = []
        for i, (_, game) in enumerate(new_games.iterrows(), 1):
            gid = game["GAME_ID"]
            print(f"  [{i}/{len(new_games)}] {game['GAME_DATE']} | "
                  f"{game['AWAY_TEAM']} @ {game['HOME_TEAM']}  (ID: {gid})")

            baskets = get_first_baskets(
                gid,
                int(game["HOME_TEAM_ID"]),
                int(game["AWAY_TEAM_ID"]),
            )

            rec = {
                "GAME_ID": gid,
                "GAME_DATE": game["GAME_DATE"],
                "HOME_TEAM": game["HOME_TEAM"],
                "AWAY_TEAM": game["AWAY_TEAM"],
            }

            if "HOME" in baskets:
                h = baskets["HOME"]
                rec["HOME_PLAYER"] = h["PLAYER_NAME"]
                rec["HOME_DESC"]   = h["DESCRIPTION"]
                rec["HOME_PTS"]    = h["POINTS"]
                rec["HOME_TIME"]   = h["PCTIMESTRING"]
                rec["HOME_PERIOD"] = h["PERIOD"]
            else:
                rec["HOME_PLAYER"] = "N/A"
                rec["HOME_DESC"]   = ""
                rec["HOME_PTS"]    = ""
                rec["HOME_TIME"]   = ""
                rec["HOME_PERIOD"] = ""

            if "AWAY" in baskets:
                a = baskets["AWAY"]
                rec["AWAY_PLAYER"] = a["PLAYER_NAME"]
                rec["AWAY_DESC"]   = a["DESCRIPTION"]
                rec["AWAY_PTS"]    = a["POINTS"]
                rec["AWAY_TIME"]   = a["PCTIMESTRING"]
                rec["AWAY_PERIOD"] = a["PERIOD"]
            else:
                rec["AWAY_PLAYER"] = "N/A"
                rec["AWAY_DESC"]   = ""
                rec["AWAY_PTS"]    = ""
                rec["AWAY_TIME"]   = ""
                rec["AWAY_PERIOD"] = ""

            new_records.append(rec)

            # Batch-write every 25 games so progress isn't lost on error
            if len(new_records) % 25 == 0:
                write_excel(new_records, output)
                new_records = []

        if new_records:
            write_excel(new_records, output)
    else:
        print("No new games to add.")

    # ── Backfill pass: fix any rows that previously had missing data ──────────
    incomplete = load_incomplete_rows(output)
    if incomplete:
        print(f"\nFound {len(incomplete)} rows with missing data — backfilling...")
        filled = 0
        for j, inc in enumerate(incomplete, 1):
            gid       = inc["game_id"]
            home_abb  = inc["home_team"]
            away_abb  = inc["away_team"]
            home_id   = TEAM_ABB_MAP.get(home_abb)
            away_id   = TEAM_ABB_MAP.get(away_abb)

            if home_id is None or away_id is None:
                print(f"  [{j}/{len(incomplete)}] Skipping {gid} — unknown team abbrev "
                      f"({home_abb}/{away_abb})")
                continue

            missing_label = "/".join(
                (["HOME"] if inc["missing_home"] else []) +
                (["AWAY"] if inc["missing_away"] else [])
            )
            print(f"  [{j}/{len(incomplete)}] {inc['game_date']} | "
                  f"{away_abb} @ {home_abb}  (ID: {gid})  — missing: {missing_label}")

            baskets = get_first_baskets(gid, int(home_id), int(away_id))
            if not baskets:
                print(f"    Still no data available — skipping.")
                continue

            patch_excel_row(
                output, inc["row_num"], baskets,
                inc["missing_home"], inc["missing_away"],
            )
            filled += 1

        print(f"  Backfill complete: {filled}/{len(incomplete)} rows updated.")
    else:
        print("\nNo incomplete rows found — dataset is fully populated.")

    print("Refreshing summary sheet...")
    add_summary_sheet(output)
    print(f"\n✅ Done! Dataset saved to: {output}")


if __name__ == "__main__":
    main()
